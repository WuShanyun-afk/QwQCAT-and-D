# -*- coding: utf-8 -*-
"""把 docx 的段落/表格完整打印出来，方便阅读内容。"""
import sys
from docx import Document


def dump(path):
    print("=" * 70)
    print("FILE:", path)
    print("=" * 70)
    doc = Document(path)
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style = p.style.name if p.style is not None else ""
        if text:
            print("[%03d][%s] %s" % (i, style, text))
        elif style and style.lower().startswith("list"):
            print("[%03d][%s] (空)" % (i, style))
    for ti, table in enumerate(doc.tables):
        print("--- TABLE %d (%dx%d) ---" % (ti, len(table.rows), len(table.columns)))
        for row in table.rows:
            cells = [c.text.replace("\n", " / ").strip() for c in row.cells]
            print(" | ".join(cells))
    print()


if __name__ == "__main__":
    for p in sys.argv[1:]:
        dump(p)
